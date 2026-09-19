"""Analizador no destructivo para el padrón histórico de membresía."""
import csv
import io
import json
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Literal, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from geo_address import normalize_address_document
from membership_documents import require_direct_membership_manager
from server import db, get_current_user


router = APIRouter(prefix="/api/membership/import", tags=["membership-import"])
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
FIELDS = [
    "first_name", "last_name", "full_name", "email", "phone", "birth_date",
    "person_number", "member_number", "address1", "address2", "city", "state",
    "zip", "household",
]
ALIASES = {
    "first_name": {"nombre", "nombres", "first name", "firstname", "first_name"},
    "last_name": {"apellido", "apellidos", "last name", "lastname", "last_name"},
    "full_name": {"nombre completo", "full name", "fullname", "miembro"},
    "email": {"email", "correo", "correo electronico", "e-mail"},
    "phone": {"telefono", "teléfono", "phone", "celular", "movil", "móvil", "whatsapp"},
    "birth_date": {"fecha nacimiento", "fecha de nacimiento", "birth date", "dob", "nacimiento"},
    "person_number": {"numero vv", "número vv", "person number", "person_number", "vv"},
    "member_number": {"numero miembro", "número miembro", "member number", "membership number", "member_number"},
    "address1": {"direccion", "dirección", "address", "address1", "calle"},
    "address2": {"direccion 2", "dirección 2", "address2", "apartamento", "unidad", "apt"},
    "city": {"ciudad", "city"}, "state": {"estado", "state", "provincia"},
    "zip": {"zip", "zipcode", "codigo postal", "código postal", "postal code"},
    "household": {"hogar", "familia", "household", "family"},
}


class DryRunResponse(BaseModel):
    dry_run: Literal[True] = True
    database_writes: Literal[0] = 0
    file: dict
    columns: list[str]
    mapping: dict[str, Optional[str]]
    summary: dict[str, int]
    rows: list[dict]
    households: list[dict]


def _key(value: object) -> str:
    raw = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", raw).strip()


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _parse_csv(content: bytes) -> tuple[list[str], list[dict]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    columns = [_text(item) for item in (reader.fieldnames or [])]
    return columns, [{column: _text(row.get(column)) for column in columns} for row in reader]


def _parse_xlsx(content: bytes) -> tuple[list[str], list[dict]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        columns = [_text(item) for item in next(iterator, [])]
        rows = [{column: _text(value) for column, value in zip(columns, values)} for values in iterator]
        workbook.close()
        return columns, rows
    except Exception as exc:
        raise HTTPException(status_code=422, detail="No se pudo leer el archivo XLSX") from exc


def _suggest_mapping(columns: list[str]) -> dict[str, Optional[str]]:
    normalized = {_key(column): column for column in columns if column}
    return {
        field: next((normalized[alias] for alias in ALIASES[field] if alias in normalized), None)
        for field in FIELDS
    }


def _mapping(value: Optional[str], columns: list[str]) -> dict[str, Optional[str]]:
    if not value:
        return _suggest_mapping(columns)
    try:
        supplied = json.loads(value)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=422, detail="El mapeo de columnas no es JSON válido") from exc
    if not isinstance(supplied, dict):
        raise HTTPException(status_code=422, detail="El mapeo de columnas debe ser un objeto")
    return {field: supplied.get(field) if supplied.get(field) in columns else None for field in FIELDS}


def _mapped(row: dict, mapping: dict, field: str) -> str:
    column = mapping.get(field)
    return _text(row.get(column)) if column else ""


def _phone(value: str) -> str:
    digits = re.sub(r"\D", "", value)
    return digits[-10:] if len(digits) >= 10 else digits


def _email(value: str) -> str:
    return value.strip().lower()


def _name_parts(row: dict, mapping: dict) -> tuple[str, str]:
    first, last = _mapped(row, mapping, "first_name"), _mapped(row, mapping, "last_name")
    if first or last:
        return first, last
    parts = _mapped(row, mapping, "full_name").split()
    return (" ".join(parts[:-1]), parts[-1]) if len(parts) > 1 else ((parts[0] if parts else ""), "")


async def _existing_indexes(rows: list[dict]) -> dict:
    emails = sorted({item["email"] for item in rows if item["email"]})
    phones = sorted({item["phone"] for item in rows if item["phone"]})
    person_numbers = sorted({item["person_number"].upper() for item in rows if item["person_number"]})
    member_numbers = sorted({item["member_number"].upper() for item in rows if item["member_number"]})
    search_keys = sorted({item["search_key"] for item in rows if item["birth_date"] and item["search_key"]})
    contacts = await db.person_contacts.find(
        {"tipo": {"$in": ["email", "telefono", "whatsapp"]}},
        {"_id": 0, "person_id": 1, "tipo": 1, "valor": 1},
    ).to_list(100000)
    matched_emails = {_email(item["valor"]): item["person_id"] for item in contacts if item["tipo"] == "email" and _email(item["valor"]) in emails}
    matched_phones = {_phone(item["valor"]): item["person_id"] for item in contacts if item["tipo"] in {"telefono", "whatsapp"} and _phone(item["valor"]) in phones}
    persons = await db.persons.find(
        {"$or": [{"person_number": {"$in": person_numbers}}, {"search_key": {"$in": search_keys}}]},
        {"_id": 0, "person_number": 1, "search_key": 1, "fecha_nacimiento": 1},
    ).to_list(100000) if person_numbers or search_keys else []
    registries = await db.membership_number_registry.find(
        {"member_number": {"$in": member_numbers}}, {"_id": 0, "member_number": 1, "person_id": 1}
    ).to_list(100000) if member_numbers else []
    return {
        "emails": matched_emails,
        "phones": matched_phones,
        "person_numbers": {str(item.get("person_number") or "").upper() for item in persons},
        "name_birth": {(item.get("search_key"), str(item.get("fecha_nacimiento") or "")[:10]) for item in persons},
        "member_numbers": {str(item["member_number"]).upper() for item in registries},
    }


@router.post("/dry-run", response_model=DryRunResponse)
async def membership_import_dry_run(
    file: UploadFile = File(...),
    mapping: Optional[str] = Form(default=None),
    current_user: dict = Depends(get_current_user),
):
    require_direct_membership_manager(current_user)
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise HTTPException(status_code=415, detail="Use un archivo CSV o XLSX")
    content = await file.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="El archivo excede 5 MiB")
    columns, source_rows = _parse_csv(content) if suffix == ".csv" else _parse_xlsx(content)
    if not columns or not any(columns):
        raise HTTPException(status_code=422, detail="El archivo no contiene encabezados")
    if len(source_rows) > MAX_ROWS:
        raise HTTPException(status_code=413, detail=f"El archivo excede {MAX_ROWS} filas")
    selected_mapping = _mapping(mapping, columns)
    normalized_rows = []
    for index, source in enumerate(source_rows, start=2):
        first, last = _name_parts(source, selected_mapping)
        address = {
            "linea1": _mapped(source, selected_mapping, "address1"),
            "linea2": _mapped(source, selected_mapping, "address2") or None,
            "ciudad": _mapped(source, selected_mapping, "city") or "Columbus",
            "provincia": _mapped(source, selected_mapping, "state") or "OH",
            "codigo_postal": _mapped(source, selected_mapping, "zip") or None,
            "pais": "US",
        }
        address_fields = normalize_address_document(address) if address["linea1"] else {}
        normalized_rows.append({
            "row_number": index, "first_name": first.strip(), "last_name": last.strip(),
            "email": _email(_mapped(source, selected_mapping, "email")),
            "phone": _phone(_mapped(source, selected_mapping, "phone")),
            "birth_date": _mapped(source, selected_mapping, "birth_date")[:10],
            "person_number": _mapped(source, selected_mapping, "person_number").upper(),
            "member_number": _mapped(source, selected_mapping, "member_number").upper(),
            "household_label": _mapped(source, selected_mapping, "household"),
            "address": address if address["linea1"] else None,
            "normalized_address_key": address_fields.get("normalized_address_key"),
            "search_key": _key(f"{first} {last}"),
        })
    existing = await _existing_indexes(normalized_rows)
    seen = {"email": {}, "phone": {}, "person_number": {}, "member_number": {}, "name_birth": {}}
    results = []
    for item in normalized_rows:
        errors, reviews, duplicates = [], [], []
        if not item["first_name"] or not item["last_name"]: errors.append("nombre_o_apellido_faltante")
        if item["email"] and not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", item["email"]): errors.append("correo_invalido")
        if item["phone"] and len(item["phone"]) < 7: reviews.append("telefono_incompleto")
        if not item["address"]: reviews.append("direccion_faltante")
        checks = {
            "email": item["email"], "phone": item["phone"], "person_number": item["person_number"],
            "member_number": item["member_number"],
            "name_birth": (item["search_key"], item["birth_date"]) if item["birth_date"] else None,
        }
        for field, value in checks.items():
            if not value: continue
            external_match = value in existing.get(f"{field}s", set()) if field not in {"name_birth"} else value in existing["name_birth"]
            if external_match: duplicates.append(f"{field}_existente")
            if value in seen[field]: duplicates.append(f"{field}_repetido_fila_{seen[field][value]}")
            else: seen[field][value] = item["row_number"]
        if duplicates: reviews.append("posible_duplicado")
        status = "error" if errors else "review" if reviews else "ready"
        results.append({**item, "status": status, "errors": errors, "review_reasons": reviews, "duplicate_matches": duplicates, "household_group": None})
    label_addresses = {}
    for item in results:
        if item["household_label"]:
            label_addresses.setdefault(_key(item["household_label"]), set()).add(item["normalized_address_key"] or f"row:{item['row_number']}")
    conflicting_labels = {label for label, addresses in label_addresses.items() if len(addresses) > 1}
    address_groups = {}
    for item in results:
        if item["normalized_address_key"]:
            address_groups.setdefault(item["normalized_address_key"], []).append(item)
        if _key(item["household_label"]) in conflicting_labels:
            item["review_reasons"].append("hogar_con_direcciones_diferentes")
            if item["status"] == "ready": item["status"] = "review"
    households = []
    for number, (address_key, members) in enumerate((entry for entry in address_groups.items() if len(entry[1]) > 1), start=1):
        group_id = f"HH-{number:03d}"
        for item in members: item["household_group"] = group_id
        households.append({"group_id": group_id, "normalized_address_key": address_key, "row_numbers": [item["row_number"] for item in members], "member_count": len(members), "action": "suggest_only"})
    summary = {"total": len(results), "ready": 0, "review": 0, "error": 0, "households": len(households)}
    for item in results: summary[item["status"]] += 1
    return DryRunResponse(
        file={"name": file.filename, "type": suffix.lstrip("."), "rows": len(results)}, columns=columns,
        mapping=selected_mapping, summary=summary, rows=results, households=households,
    )